"""Build the per-fair share directory and the hub-wide master tracker."""
import math
import re
from pathlib import Path

import segno
from PIL import Image, ImageOps
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .config import (APPLIED_VIA, CATEGORIES, CATEGORY, DATE_COLUMNS, IN_PROGRESS, MASTER_PATH, SHARE_CATEGORIES,
                     STATUS_FILL, STATUSES, TRACK_COLUMNS, TRACKABLE_CATEGORIES)
from .fairs import list_fairs, load_companies, load_json, load_meta

QR_PX = 112
LASTROW_REF = 1000
_thin = Side(style="thin", color="D9D9D9")
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
FONT = Font(name="Calibri", size=10)
LINK_FONT = Font(name="Calibri", size=10, color="0563C1", underline="single")
CAT_ORDER = {k: i for i, (k, *_) in enumerate(CATEGORIES)}
TRACK_NAMES = [n for n, _ in TRACK_COLUMNS]


# ---------------------------------------------------------------- helpers
def _slug(s):
    return re.sub(r"[^A-Za-z0-9]+", "_", s).strip("_")[:50]


def _dash(v):
    if isinstance(v, list):
        v = "\n".join(v)
    return v if v not in (None, "") else "-"


def qr_image(fair_dir, c):
    gen = Path(fair_dir) / "work" / "qr" / "generated"
    gen.mkdir(parents=True, exist_ok=True)
    if c.get("qr_photo"):
        qp = c["qr_photo"]
        img = ImageOps.exif_transpose(Image.open(Path(fair_dir) / qp["file"])).convert("RGB")
        W, H = img.size
        x0, y0, x1, y1 = qp["box"]
        crop = img.crop((int(x0 * W), int(y0 * H), int(x1 * W), int(y1 * H)))
        crop.thumbnail((700, 700))
        p = gen / f"{_slug(c['name'])}_photo.png"
        crop.save(p)
        return p
    if c["links"]:
        p = gen / f"{_slug(c['name'])}.png"
        segno.make(c["links"][0]["url"], error="m").save(p, scale=6, border=2)
        return p
    return None


def display_link(url):
    if url.startswith("mailto:"):
        return f"Email → {url[7:].split('?')[0]}" + (" (pre-filled)" if "?" in url else "")
    return url


def primary_link_text(c):
    if c.get("qr_photo"):
        return c["qr_photo"]["label"]
    if c["links"]:
        return f"{c['links'][0]['label']}\n{display_link(c['links'][0]['url'])}"
    return "-"


def other_links_text(c):
    links = c["links"] if c.get("qr_photo") else c["links"][1:]
    return "\n".join(f"{l['label']}: {display_link(l['url'])}" for l in links) or "-"


def link_target(c):
    if c.get("qr_photo") and c["email"]:
        return "mailto:" + c["email"].split()[0]
    return c["links"][0]["url"] if c["links"] else None


def est_height(values_widths, has_img):
    lines = 1
    for text, width in values_widths:
        cpl = max(8, int(width * 1.15))
        lines = max(lines, sum(max(1, math.ceil(len(p) / cpl)) for p in str(text).split("\n")))
    h = lines * 13.2 + 6
    if has_img:
        h = max(h, QR_PX * 0.75 + 10)
    return min(h, 409)


def cf_fill(color):
    return PatternFill(start_color=color, end_color=color, fill_type="solid")


def style_header(cell, fill):
    cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = BORDER


def fair_title(meta):
    dates = meta.get("dates") or []
    when = f", {dates[0]}" + (f" to {dates[-1]}" if len(dates) > 1 else "") if dates else ""
    return f"{meta.get('short_name', meta['id'])} Career Fair ({meta.get('venue', '')}{when})"


# Company-info columns: (header, width, value function)
INFO = {
    "Booth": (7, lambda c: _dash(c["booth"])),
    "Sponsor Tier": (14, lambda c: _dash(c["tier"])),
    "Industry": (22, lambda c: _dash(c["industry"])),
    "Programme Type": (20, lambda c: _dash(c["programme_type"])),
    "IT / CS Roles": (24, lambda c: _dash(c["it_roles"])),
    "Roles / Positions / Programmes": (60, lambda c: _dash(c["roles"])),
    "Requirements / Eligibility": (38, lambda c: _dash(c["requirements"])),
    "Min CGPA": (12, lambda c: _dash(c["min_cgpa"])),
    "Intake / Timing": (20, lambda c: _dash(c["intake"])),
    "About the Company": (44, lambda c: _dash(c["about"])),
    "Contact Person (hiring contact)": (26, lambda c: _dash(c["contact"])),
    "Email": (28, lambda c: _dash(c["email"])),
    "Phone / WhatsApp": (18, lambda c: _dash(c["phone"])),
    "Website": (26, lambda c: _dash(c["website"])),
    "LinkedIn / Social": (26, lambda c: _dash(c["social"])),
    "How to Apply": (34, lambda c: _dash(c["how_to_apply"])),
    "Accepts Resume Drop?": (20, lambda c: _dash(c["resume_drop"])),
    "QR Code (scan)": (18, lambda c: ""),
    "QR Link (click to open)": (38, primary_link_text),
    "Other Links": (42, other_links_text),
    "Location": (26, lambda c: _dash(c["location"])),
    "In Share Copy?": (26, lambda c: "Yes" if c["share"] and c["category"] in SHARE_CATEGORIES
                       else "No — " + (c["share_reason"] or "not an application target")),
    "Source Files (photos / booklet pages)": (28, lambda c: ", ".join(c["sources"]) or "-"),
    "Extraction Notes (internal)": (48, lambda c: _dash(c["notes"])),
}
SHARE_INFO = ["Industry", "Programme Type", "Roles / Positions / Programmes", "Requirements / Eligibility",
              "About the Company", "Contact Person (hiring contact)", "Email", "Phone / WhatsApp", "Website",
              "LinkedIn / Social", "How to Apply", "QR Code (scan)", "QR Link (click to open)", "Other Links", "Location"]
MASTER_INFO = list(INFO.keys())


def _write_company(ws, row, cols, values, fair_dir, c):
    """cols: list of (header, width). values: header -> value. Adds QR image + hyperlink; sets row height."""
    col = {h: i + 1 for i, (h, _) in enumerate(cols)}
    vw = []
    for h, w in cols:
        v = values.get(h, "")
        cell = ws.cell(row=row, column=col[h], value=v)
        cell.font = FONT
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.border = BORDER
        if h not in TRACK_NAMES:
            vw.append((v, w))
    ws.cell(row=row, column=1).font = Font(name="Calibri", size=10, bold=True)
    target = link_target(c)
    if target and "QR Link (click to open)" in col:
        lc = ws.cell(row=row, column=col["QR Link (click to open)"])
        lc.hyperlink = target
        lc.font = LINK_FONT
    img_path = qr_image(fair_dir, c) if "QR Code (scan)" in col else None
    if img_path:
        img = XLImage(str(img_path))
        img.width = img.height = QR_PX
        ws.add_image(img, f"{get_column_letter(col['QR Code (scan)'])}{row}")
    ws.row_dimensions[row].height = est_height(vw, bool(img_path))


def _save(wb, path):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(path)
        return path
    except PermissionError:
        alt = path.with_name(path.stem + "_new" + path.suffix)
        wb.save(alt)
        print(f"! {path.name} is open in Excel — saved to {alt.name} instead. Close Excel and rebuild to replace it.")
        return alt


# ---------------------------------------------------------------- share copy
def build_share(fair_dir):
    fair_dir = Path(fair_dir)
    meta = load_meta(fair_dir)
    companies = [c for c in load_companies(fair_dir) if c["share"] and c["category"] in SHARE_CATEGORIES]
    cols = [("Company", 26)] + [(h, INFO[h][0]) for h in SHARE_INFO]
    ncol, last = len(cols), get_column_letter(len(cols))

    wb = Workbook()
    ws = wb.active
    ws.title = f"{meta.get('short_name', 'Fair')} Directory"[:31]
    ws["A1"] = f"{fair_title(meta)} — Job, Internship & Programme Directory"
    ws["A1"].font = Font(size=16, bold=True, color="1F4E78")
    checked = meta.get("links_checked") or "—"
    ws["A2"] = ("Compiled from booth photos and the official programme booklet. QR codes were scanned and "
                f"re-generated so you can scan them from this sheet; the QR Link column is clickable. Links checked {checked}.")
    ws["A2"].font = Font(size=10, italic=True, color="595959")
    ws.merge_cells(f"A1:{last}1")
    ws.merge_cells(f"A2:{last}2")
    ws.row_dimensions[1].height = 24
    for i, (h, w) in enumerate(cols, start=1):
        style_header(ws.cell(row=3, column=i, value=h), "203864")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[3].height = 32

    row = 4
    for key, label, dark, light in CATEGORIES:
        group = sorted([c for c in companies if c["category"] == key], key=lambda c: c["name"].lower())
        if not group:
            continue
        ws.cell(row=row, column=1, value=f"{label.upper()}  ({len(group)})")
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncol)
        sc = ws.cell(row=row, column=1)
        sc.font = Font(size=12, bold=True, color="FFFFFF")
        sc.fill = PatternFill("solid", fgColor=dark)
        ws.row_dimensions[row].height = 22
        row += 1
        for n, c in enumerate(group):
            values = {"Company": c["name"], **{h: INFO[h][1](c) for h in SHARE_INFO}}
            _write_company(ws, row, cols, values, fair_dir, c)
            if n % 2:
                for i in range(1, ncol + 1):
                    ws.cell(row=row, column=i).fill = PatternFill("solid", fgColor="F2F6FC")
            row += 1
    ws.freeze_panes = "B4"
    ws.auto_filter.ref = f"A3:{last}{row - 1}"
    ws.sheet_view.zoomScale = 90
    out = fair_dir / "output" / f"{meta.get('short_name', meta['id'])}_Job_Directory_Share.xlsx"
    out = _save(wb, out)
    print(f"share copy: {len(companies)} companies -> {out}")
    return out


# ---------------------------------------------------------------- master
def read_tracker(path, default_fair=None):
    """Return {(fair, company): {"tracker": {...}, "row": {...}}} from an existing master workbook."""
    path = Path(path)
    if not path.exists():
        return {}
    wb = load_workbook(path)
    if "Master Directory" not in wb.sheetnames:
        return {}
    ws = wb["Master Directory"]
    hdr_row = next(r for r in range(1, 8) if ws.cell(row=r, column=1).value == "Company")
    headers = {ws.cell(row=hdr_row, column=i).value: i for i in range(1, ws.max_column + 1)
               if ws.cell(row=hdr_row, column=i).value}
    out = {}
    for r in range(hdr_row + 1, ws.max_row + 1):
        company = ws.cell(row=r, column=1).value
        if not company:
            continue
        fair = ws.cell(row=r, column=headers["Fair"]).value if "Fair" in headers else default_fair
        row = {h: ws.cell(row=r, column=i).value for h, i in headers.items()}
        out[(fair, company)] = {"tracker": {h: row.get(h) for h in TRACK_NAMES if h in row}, "row": row}
    return out


def build_master(seed=None, seed_fair=None):
    fairs = [(d, load_meta(d), load_companies(d)) for d in list_fairs()]
    prev = read_tracker(MASTER_PATH) if MASTER_PATH.exists() else read_tracker(seed, seed_fair) if seed else {}

    wb = Workbook()
    dash = wb.active
    dash.title = "Dashboard"
    ws = wb.create_sheet("Master Directory")
    cols = [("Company", 28), ("Fair", 10), ("Category", 18)] + TRACK_COLUMNS + [(h, INFO[h][0]) for h in MASTER_INFO]
    col = {h: i + 1 for i, (h, _) in enumerate(cols)}
    L = {h: get_column_letter(i) for h, i in col.items()}
    ncol = len(cols)

    ws["A1"] = "Career Fair Hub — Master Directory & Application Tracker"
    ws["A1"].font = Font(size=16, bold=True, color="1F4E78")
    ws["A2"] = ("Red columns = your tracker (kept when the file is rebuilt). Follow-ups turn amber within 7 days and red "
                "when overdue. Filter by Fair / Category / Application Status. Collapse the tracker with [–] above column D.")
    ws["A2"].font = Font(size=10, italic=True, color="595959")
    ws.row_dimensions[1].height = 24
    for h, w in cols:
        style_header(ws.cell(row=3, column=col[h], value=h), "C00000" if h in TRACK_NAMES else "203864")
        ws.column_dimensions[L[h]].width = w
    ws.row_dimensions[3].height = 46

    row, FIRST, used = 4, 4, set()
    for fair_dir, meta, companies in fairs:
        short = meta.get("short_name", meta["id"])
        for c in sorted(companies, key=lambda c: (CAT_ORDER.get(c["category"], 99), c["name"].lower())):
            label, dark, light = CATEGORY[c["category"]]
            values = {"Company": c["name"], "Fair": short, "Category": label,
                      **{h: INFO[h][1](c) for h in MASTER_INFO}}
            if c["category"] in TRACKABLE_CATEGORIES:
                values["Application Status"] = "Not Applied"
            old = prev.get((short, c["name"]))
            if old:
                used.add((short, c["name"]))
                values.update({h: v for h, v in old["tracker"].items() if v not in (None, "")})
            _write_company(ws, row, cols, values, fair_dir, c)
            cc = ws.cell(row=row, column=col["Category"])
            cc.fill = PatternFill("solid", fgColor=light)
            cc.font = Font(size=10, bold=True, color=dark)
            row += 1
    manual = [v for k, v in prev.items() if k not in used]
    for m in manual:  # rows the user added by hand (or companies removed from data) — never drop them
        for h, i in col.items():
            cell = ws.cell(row=row, column=i, value=m["row"].get(h))
            cell.font, cell.border = FONT, BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        if not m["row"].get("Category"):
            ws.cell(row=row, column=col["Category"], value="Manual entry")
        row += 1
    LAST = row - 1
    for r in range(FIRST, LAST + 201):  # room for rows added by hand
        for dcol in DATE_COLUMNS:
            ws.cell(row=r, column=col[dcol]).number_format = "dd-mmm-yyyy"

    ws.freeze_panes = f"D{FIRST}"
    ws.auto_filter.ref = f"A3:{get_column_letter(ncol)}{LAST}"
    ws.column_dimensions.group(L[TRACK_NAMES[0]], L[TRACK_NAMES[-1]], outline_level=1, hidden=False)
    ws.sheet_view.zoomScale = 90

    def rng(h):
        return f"{L[h]}{FIRST}:{L[h]}{LASTROW_REF}"

    for h, opts in (("Interested?", ["Yes", "Maybe", "No"]), ("Priority", ["High", "Medium", "Low"]),
                    ("Application Status", STATUSES), ("Applied Via", APPLIED_VIA)):
        dv = DataValidation(type="list", formula1='"' + ",".join(opts) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(rng(h))
    for h in DATE_COLUMNS:
        dv = DataValidation(type="date", operator="greaterThan", formula1="DATE(2020,1,1)", allow_blank=True,
                            showErrorMessage=True, errorTitle="Date needed", error="Enter a date, e.g. 30/9/2026")
        ws.add_data_validation(dv)
        dv.add(rng(h))
    for status, fill in STATUS_FILL.items():
        ws.conditional_formatting.add(rng("Application Status"),
                                      CellIsRule(operator="equal", formula=[f'"{status}"'], fill=cf_fill(fill)))
    S, F, K = L["Application Status"], L["Follow-up Date"], L["Interview Date"]
    active = "OR(" + ",".join(f'${S}{FIRST}="{s}"' for s in IN_PROGRESS) + ")"
    ws.conditional_formatting.add(rng("Follow-up Date"), FormulaRule(
        formula=[f'AND(${F}{FIRST}<>"",${F}{FIRST}<TODAY(),{active})'],
        fill=cf_fill("FF7C80"), font=Font(bold=True, color="9C0006")))
    ws.conditional_formatting.add(rng("Follow-up Date"), FormulaRule(
        formula=[f'AND(${F}{FIRST}<>"",${F}{FIRST}>=TODAY(),${F}{FIRST}<=TODAY()+7,{active})'], fill=cf_fill("FFD966")))
    ws.conditional_formatting.add(rng("Interview Date"), FormulaRule(
        formula=[f'AND(${K}{FIRST}<>"",${K}{FIRST}>=TODAY())'], fill=cf_fill("C6EFCE")))
    ws.conditional_formatting.add(rng("Priority"), CellIsRule(operator="equal", formula=['"High"'],
                                                              font=Font(bold=True, color="C00000")))

    _dashboard(dash, L, FIRST, [m for _, m, _ in fairs])
    _qr_sheet(wb.create_sheet("All Scanned QR Codes"), fairs)
    _source_sheet(wb.create_sheet("Source Reference (internal)"), fairs)
    wb.active = 0
    out = _save(wb, MASTER_PATH)
    n = sum(len(c) for _, _, c in fairs)
    print(f"master: {n} companies across {len(fairs)} fair(s), {len(used)} tracker rows carried over, "
          f"{len(manual)} manual rows kept -> {out}")
    return out


def _dashboard(dash, L, FIRST, metas):
    MD = "'Master Directory'!"

    def R(h):
        return f"{MD}${L[h]}${FIRST}:${L[h]}${LASTROW_REF}"

    dash["A1"] = "Career Fair Hub — Dashboard"
    dash["A1"].font = Font(size=18, bold=True, color="1F4E78")
    dash["A2"] = "All numbers update automatically from the Master Directory sheet."
    dash["A2"].font = Font(italic=True, color="595959")
    for c, w in (("A", 44), ("B", 12), ("D", 44), ("E", 12)):
        dash.column_dimensions[c].width = w

    def block(r, c0, title, items, color):
        for i, v in ((0, title), (1, None)):
            h = dash.cell(row=r, column=c0 + i, value=v)
            h.fill = PatternFill("solid", fgColor=color)
            h.font = Font(bold=True, color="FFFFFF")
        for i, (label, formula) in enumerate(items, start=1):
            a = dash.cell(row=r + i, column=c0, value=label)
            b = dash.cell(row=r + i, column=c0 + 1, value=formula)
            b.font = Font(bold=True)
            b.alignment = Alignment(horizontal="center")
            a.border = b.border = BORDER
        return r + len(items) + 2

    ST, FU, IV = R("Application Status"), R("Follow-up Date"), R("Interview Date")
    sent = ["Applied", "Assessment / Test", "Interview", "Offer", "Accepted", "Rejected", "Withdrawn"]
    left = block(4, 1, "APPLICATION PIPELINE",
                 [("Applications sent (all stages)", "=" + "+".join(f'COUNTIF({ST},"{s}")' for s in sent)),
                  ("Still in progress (Applied / Test / Interview)",
                   "=" + "+".join(f'COUNTIF({ST},"{s}")' for s in ("Applied", "Assessment / Test", "Interview")))] +
                 [(s, f'=COUNTIF({ST},"{s}")') for s in STATUSES], "C00000")
    in_prog = "+".join(f'({ST}="{s}")' for s in IN_PROGRESS)
    left = block(left, 1, "FOLLOW-UPS & INTERVIEWS",
                 [("Follow-ups OVERDUE", f'=SUMPRODUCT(({FU}<>"")*({FU}<TODAY())*(({in_prog})>0))'),
                  ("Follow-ups due in the next 7 days", f'=COUNTIFS({FU},">="&TODAY(),{FU},"<="&TODAY()+7)'),
                  ("Upcoming interviews", f'=COUNTIFS({IV},">="&TODAY())')], "BF8F00")
    block(left, 1, "INTEREST & PRIORITY",
          [(f"Interested — {v}", f'=COUNTIF({R("Interested?")},"{v}")') for v in ("Yes", "Maybe")] +
          [(f"Priority — {v}", f'=COUNTIF({R("Priority")},"{v}")') for v in ("High", "Medium", "Low")], "7030A0")

    right = block(4, 4, "COMPANIES BY FAIR",
                  [(m.get("short_name", m["id"]), f'=COUNTIF({R("Fair")},"{m.get("short_name", m["id"])}")') for m in metas] +
                  [("Total", f'=COUNTA({R("Company")})')], "203864")
    right = block(right, 4, "COMPANIES BY CATEGORY",
                  [(label, f'=COUNTIF({R("Category")},"{label}")') for _, label, _, _ in CATEGORIES], "1F4E78")
    PT, IT = R("Programme Type"), R("IT / CS Roles")
    right = block(right, 4, "OPPORTUNITY MIX (job booths)",
                  [("Offer internships", f'=COUNTIF({PT},"*Internship*")'),
                   ("Graduate / rotation / trainee programmes",
                    f'=COUNTIF({PT},"*Graduate*")+COUNTIF({PT},"*Trainee*")-COUNTIFS({PT},"*Graduate*",{PT},"*Trainee*")'),
                   ("Direct hire / open roles", f'=COUNTIF({PT},"*Direct hire*")'),
                   ("Apprenticeship", f'=COUNTIF({PT},"*Apprentice*")'),
                   ("List IT / CS roles", f'=COUNTIF({IT},"Yes*")'),
                   ("Accept resume / CV directly", f'=COUNTIF({R("Accepts Resume Drop?")},"Yes*")')], "548235")
    tips = dash.cell(row=right, column=4, value=(
        "Tips:\n• Set 'Application Status' when you apply, and put a Follow-up Date ~7–10 days later.\n"
        "• Filter Category = 'Job / Internship Booth' and IT / CS Roles = 'Yes…' for tech roles.\n"
        "• 'Other Exhibitor' rows are real employers with no apply details at the fair — use their careers sites.\n"
        "• Your tracker columns survive `python -m cfhub build`; add your own rows at the bottom if you like."))
    tips.alignment = Alignment(wrap_text=True, vertical="top")
    dash.merge_cells(start_row=right, start_column=4, end_row=right + 11, end_column=5)


def _qr_sheet(qs, fairs):
    cols = [("#", 5), ("Fair", 10), ("Source File", 16), ("From", 12), ("Company", 30), ("What it is / Status", 40),
            ("Decoded Link / Content", 70), ("QR (from photo)", 14)]
    for i, (h, w) in enumerate(cols, start=1):
        style_header(qs.cell(row=1, column=i, value=h), "203864")
        qs.column_dimensions[get_column_letter(i)].width = w
    r = 2
    for fair_dir, meta, _ in fairs:
        thumbs = Path(fair_dir) / "work" / "qr" / "thumbs"
        thumbs.mkdir(parents=True, exist_ok=True)
        for e in load_json(fair_dir, "work/qr/qr_index.json", []):
            text = e.get("text") or "-"
            vals = [r - 1, meta.get("short_name", meta["id"]), e["source"],
                    "Booth photo" if e["source"].startswith("IMG") else "Booklet", e.get("company") or "-",
                    e.get("status") or "-", display_link(text) if text.startswith("mailto:") else text]
            for i, v in enumerate(vals, start=1):
                cell = qs.cell(row=r, column=i, value=v)
                cell.font, cell.border = FONT, BORDER
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            if text.startswith(("http", "mailto:")):
                qs.cell(row=r, column=7).hyperlink = text
                qs.cell(row=r, column=7).font = LINK_FONT
            if any(k in (e.get("status") or "") for k in ("DEAD", "Could not", "review")):
                qs.cell(row=r, column=6).font = Font(size=10, bold=True, color="C00000")
            crop = Path(fair_dir) / e["crop"] if e.get("crop") else None
            if crop and crop.exists():
                t = thumbs / f"{r:04d}.png"
                im = Image.open(crop).convert("RGB")
                im.thumbnail((160, 160))
                im.save(t)
                xi = XLImage(str(t))
                xi.width = xi.height = 80
                qs.add_image(xi, f"H{r}")
            qs.row_dimensions[r].height = 64
            r += 1
    qs.freeze_panes = "A2"
    qs.auto_filter.ref = f"A1:G{max(r - 1, 1)}"


def _source_sheet(sr, fairs):
    keys = ["batch", "source_file", "name", "booth", "booth_type", "industry", "roles", "requirements", "description",
            "contact", "email", "phone", "website", "social", "qr_code", "resume_drop", "application_instructions",
            "location", "notes"]
    sr.append(["fair"] + keys)
    for cell in sr[1]:
        style_header(cell, "595959")
    for fair_dir, meta, _ in fairs:
        for rec in load_json(fair_dir, "work/extraction/raw_records.json", []):
            sr.append([meta.get("short_name", meta["id"])] + [rec.get(k, "") if not isinstance(rec.get(k), list)
                                                               else ", ".join(rec[k]) for k in keys])
    for i in range(1, len(keys) + 2):
        sr.column_dimensions[get_column_letter(i)].width = 24
    for row in sr.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    sr.freeze_panes = "D2"
    sr.auto_filter.ref = sr.dimensions
